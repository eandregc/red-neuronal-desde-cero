"""
Exportación de los pesos a Excel y lectura de vuelta.

El libro no solo guarda los números: las hojas de comprobación vuelven a
calcular la red con fórmulas de Excel, de modo que se puede verificar a mano
que los pesos producen las salidas reportadas por Python.

    Pesos             configuración del entrenamiento y pesos finales
    Comprobación      propagación hacia adelante con fórmulas (ŷ Excel vs ŷ Python)
    Retropropagación  primera época de entrenamiento rehecha con fórmulas
    Pesos iniciales   pesos aleatorios con los que arrancó el entrenamiento
    Entrenamiento     error cuadrático medio por época y su gráfica
"""

from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ARCHIVO_PESOS = Path(__file__).resolve().parent.parent / "pesos.xlsx"

FUENTE = "Arial"
AZUL = "0000FF"   # valores generados por Python (datos y pesos)
VERDE = "008000"  # referencias a otra hoja
NEGRO = "000000"  # fórmulas y etiquetas
RELLENO_TITULO = PatternFill("solid", fgColor="1F4E78")
RELLENO_GRUPO = PatternFill("solid", fgColor="D9E1F2")
FMT_PESO = "0.000000"
FMT_DIF = "0.00E+00"

FILA_CONFIG = 4   # la tabla de configuración empieza aquí en la hoja Pesos
FILA_PESOS = 15   # y los bloques de pesos aquí (leer_pesos depende de ambas)


# --------------------------------------------------------------------------- utilidades

def _celda(ws, fila, col, valor, color=NEGRO, negrita=False, formato=None):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = Font(name=FUENTE, color=color, bold=negrita)
    if formato:
        c.number_format = formato
    return c


def _titulo(ws, texto, subtitulo=None):
    ws.cell(row=1, column=1, value=texto).font = Font(name=FUENTE, size=14, bold=True)
    if subtitulo:
        ws.cell(row=2, column=1, value=subtitulo).font = Font(name=FUENTE, italic=True, color="595959")


def _encabezados(ws, fila, textos, col=1):
    for k, texto in enumerate(textos):
        c = _celda(ws, fila, col + k, texto, negrita=True)
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.fill = RELLENO_TITULO
        c.alignment = Alignment(horizontal="center")


def _grupo(ws, fila, col_ini, col_fin, texto):
    """Encabezado agrupado (celdas combinadas) sobre varias columnas."""
    if col_fin > col_ini:
        ws.merge_cells(start_row=fila, start_column=col_ini, end_row=fila, end_column=col_fin)
    c = ws.cell(row=fila, column=col_ini, value=texto)
    c.font = Font(name=FUENTE, bold=True)
    c.fill = RELLENO_GRUPO
    c.alignment = Alignment(horizontal="center")


def _ref(ws, fila, col):
    """Referencia absoluta con nombre de hoja, p. ej. 'Pesos'!$B$16."""
    return f"'{ws.title}'!${get_column_letter(col)}${fila}"


def _col(k):
    return get_column_letter(k)


def _anchos(ws, anchos):
    for k, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[_col(k)].width = ancho


def _leyenda(ws, fila):
    _celda(ws, fila, 1, "Valores generados por Python", AZUL)
    _celda(ws, fila + 1, 1, "Fórmulas de Excel", NEGRO)
    _celda(ws, fila + 2, 1, "Referencias a otra hoja", VERDE)


def _escribir_pesos(ws, fila, pesos):
    """Escribe W1, b1, W2 y b2 a partir de 'fila' y regresa la referencia de cada celda."""
    n_ent, n_ocu = pesos["W1"].shape
    refs = {"W1": [[None] * n_ocu for _ in range(n_ent)], "b1": [], "W2": [], "b2": None}

    _celda(ws, fila, 1, "Capa oculta: W1 (entrada → neurona oculta) y sesgos b1", negrita=True)
    _encabezados(ws, fila + 1, [""] + [f"h{j + 1}" for j in range(n_ocu)])
    for i in range(n_ent):
        _celda(ws, fila + 2 + i, 1, f"x{i + 1}", negrita=True)
        for j in range(n_ocu):
            _celda(ws, fila + 2 + i, 2 + j, float(pesos["W1"][i, j]), AZUL, formato=FMT_PESO)
            refs["W1"][i][j] = _ref(ws, fila + 2 + i, 2 + j)
    f = fila + 2 + n_ent
    _celda(ws, f, 1, "b1 (sesgo)", negrita=True)
    for j in range(n_ocu):
        _celda(ws, f, 2 + j, float(pesos["b1"][0, j]), AZUL, formato=FMT_PESO)
        refs["b1"].append(_ref(ws, f, 2 + j))

    f += 2
    _celda(ws, f, 1, "Capa de salida: W2 (neurona oculta → salida) y sesgo b2", negrita=True)
    _encabezados(ws, f + 1, ["", "ŷ"])
    for j in range(n_ocu):
        _celda(ws, f + 2 + j, 1, f"h{j + 1}", negrita=True)
        _celda(ws, f + 2 + j, 2, float(pesos["W2"][j, 0]), AZUL, formato=FMT_PESO)
        refs["W2"].append(_ref(ws, f + 2 + j, 2))
    f += 2 + n_ocu
    _celda(ws, f, 1, "b2 (sesgo)", negrita=True)
    _celda(ws, f, 2, float(pesos["b2"][0, 0]), AZUL, formato=FMT_PESO)
    refs["b2"] = _ref(ws, f, 2)
    return refs


def _propagacion(ws, fila, X, y, refs):
    """
    Escribe la tabla de propagación hacia adelante con fórmulas (una fila por muestra).
    Regresa un diccionario con la columna donde quedó cada cantidad.
    """
    n_ent, n_ocu = len(refs["W1"]), len(refs["b1"])
    c = {"x": [1 + i for i in range(n_ent)], "y": 1 + n_ent}
    c["z"] = [c["y"] + 1 + j for j in range(n_ocu)]
    c["a"] = [c["z"][-1] + 1 + j for j in range(n_ocu)]
    c["z_sal"] = c["a"][-1] + 1
    c["y_hat"] = c["z_sal"] + 1

    _grupo(ws, fila, 1, c["y"], "Datos")
    _grupo(ws, fila, c["z"][0], c["z"][-1], "z = x·W1 + b1")
    _grupo(ws, fila, c["a"][0], c["a"][-1], "a = σ(z)")
    _grupo(ws, fila, c["z_sal"], c["y_hat"], "Salida")
    _encabezados(ws, fila + 1,
                 [f"x{i + 1}" for i in range(n_ent)] + ["y"]
                 + [f"z_h{j + 1}" for j in range(n_ocu)]
                 + [f"a_h{j + 1}" for j in range(n_ocu)]
                 + ["z_salida", "ŷ (Excel)"])

    for k in range(len(X)):
        r = fila + 2 + k
        for i in range(n_ent):
            _celda(ws, r, c["x"][i], float(X[k, i]), AZUL)
        _celda(ws, r, c["y"], float(y[k, 0]), AZUL)
        for j in range(n_ocu):
            suma = "+".join(f"{_col(c['x'][i])}{r}*{refs['W1'][i][j]}" for i in range(n_ent))
            _celda(ws, r, c["z"][j], f"={suma}+{refs['b1'][j]}", formato=FMT_PESO)
            _celda(ws, r, c["a"][j], f"=1/(1+EXP(-{_col(c['z'][j])}{r}))", formato=FMT_PESO)
        suma = "+".join(f"{_col(c['a'][j])}{r}*{refs['W2'][j]}" for j in range(n_ocu))
        _celda(ws, r, c["z_sal"], f"={suma}+{refs['b2']}", formato=FMT_PESO)
        _celda(ws, r, c["y_hat"], f"=1/(1+EXP(-{_col(c['z_sal'])}{r}))", formato=FMT_PESO)

    c["filas"] = (fila + 2, fila + 1 + len(X))
    return c


def _rango(col, filas):
    return f"${_col(col)}${filas[0]}:${_col(col)}${filas[1]}"


# --------------------------------------------------------------------------- hojas

def _hoja_pesos(ws, config, pesos, error_final):
    _titulo(ws, "Red neuronal desde cero: pesos generados",
            "Pesos finales del entrenamiento. Las demás hojas los usan con fórmulas para comprobarlos.")
    n_ent, n_ocu = pesos["W1"].shape
    col_fin = 1 + max(n_ocu, 3)  # el valor se combina sobre las columnas de los pesos
    _encabezados(ws, FILA_CONFIG - 1, ["Parámetro", "Valor"])
    ws.merge_cells(start_row=FILA_CONFIG - 1, start_column=2, end_row=FILA_CONFIG - 1, end_column=col_fin)
    filas = [
        ("Compuerta", config["compuerta"].upper()),
        ("Arquitectura", config["arquitectura"]),
        ("Neuronas ocultas", config["ocultas"]),
        ("Función de activación", "Sigmoide σ(z) = 1 / (1 + e^(−z))"),
        ("Función de error", "ECM = (1/n) Σ (ŷ − y)²"),
        ("Épocas", config["epocas"]),
        ("Tasa de aprendizaje (η)", config["tasa"]),
        ("Semilla aleatoria", config["semilla"]),
        ("ECM final (Python)", error_final),
    ]
    celdas = {}
    for k, (nombre, valor) in enumerate(filas):
        f = FILA_CONFIG + k
        _celda(ws, f, 1, nombre, negrita=True)
        c = _celda(ws, f, 2, valor, AZUL, formato=FMT_DIF if "ECM final" in nombre else None)
        c.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=f, start_column=2, end_row=f, end_column=col_fin)
        celdas[nombre] = _ref(ws, f, 2)

    refs = _escribir_pesos(ws, FILA_PESOS, pesos)
    _leyenda(ws, FILA_PESOS + 6 + n_ent + n_ocu + 2)  # dos filas debajo de b2
    _anchos(ws, [26] + [12] * (col_fin - 1))
    return refs, celdas


def _hoja_comprobacion(ws, X, y, refs, y_pred, celdas_config):
    _titulo(ws, "Comprobación: propagación hacia adelante",
            "Cada celda se calcula con fórmulas a partir de la hoja Pesos. "
            "ŷ (Excel) debe coincidir con ŷ (Python).")
    c = _propagacion(ws, 4, X, y, refs)
    f0, f1 = c["filas"]
    c_py, c_dif, c_clase, c_ok = c["y_hat"] + 1, c["y_hat"] + 2, c["y_hat"] + 3, c["y_hat"] + 4

    _grupo(ws, 4, c_py, c_ok, "Comprobación")
    _encabezados(ws, 5, ["ŷ (Python)", "|Diferencia|", "Clase", "¿Correcto?"], col=c_py)
    for k in range(len(X)):
        r = f0 + k
        yh = f"{_col(c['y_hat'])}{r}"
        _celda(ws, r, c_py, float(y_pred[k, 0]), AZUL, formato=FMT_PESO)
        _celda(ws, r, c_dif, f"=ABS({yh}-{_col(c_py)}{r})", formato=FMT_DIF)
        _celda(ws, r, c_clase, f"=IF({yh}>=0.5,1,0)").alignment = Alignment(horizontal="center")
        _celda(ws, r, c_ok, f'=IF({_col(c_clase)}{r}={_col(c["y"])}{r},"Sí","No")'
               ).alignment = Alignment(horizontal="center")

    r = f1 + 2
    resumen = [
        ("ECM (Excel)", f"=SUMXMY2({_rango(c['y_hat'], c['filas'])},{_rango(c['y'], c['filas'])})"
                        f"/COUNT({_rango(c['y'], c['filas'])})", NEGRO, FMT_DIF),
        ("ECM (Python)", f"={celdas_config['ECM final (Python)']}", VERDE, FMT_DIF),
        ("Diferencia máxima Excel vs Python", f"=MAX({_rango(c_dif, c['filas'])})", NEGRO, FMT_DIF),
        ("Exactitud", f'=COUNTIF({_rango(c_ok, c["filas"])},"Sí")/COUNT({_rango(c["y"], c["filas"])})',
         NEGRO, "0%"),
    ]
    for k, (nombre, formula, color, formato) in enumerate(resumen):
        _celda(ws, r + k, 1, nombre, negrita=True)
        _celda(ws, r + k, 6, formula, color, formato=formato)
    _leyenda(ws, r + len(resumen) + 1)
    _anchos(ws, [9] * len(c["x"]) + [9] + [12] * (c_ok - c["y"]))
    ws.freeze_panes = "A6"


def _hoja_retropropagacion(ws, X, y, refs_ini, pesos_1, celdas_config):
    n_ent, n_ocu = len(refs_ini["W1"]), len(refs_ini["b1"])
    _titulo(ws, "Comprobación: retropropagación (primera época)",
            "Parte de los pesos iniciales, calcula los gradientes con la regla de la cadena "
            "y aplica un paso de descenso de gradiente con fórmulas.")
    _celda(ws, 4, 1, "Tasa de aprendizaje (η)", negrita=True)
    _celda(ws, 4, 3, f"={celdas_config['Tasa de aprendizaje (η)']}", VERDE)
    _celda(ws, 5, 1, "n (muestras)", negrita=True)
    eta, n = "$C$4", "$C$5"

    c = _propagacion(ws, 7, X, y, refs_ini)
    filas = c["filas"]
    _celda(ws, 5, 3, f"=COUNT({_rango(c['y'], filas)})")
    c_e, c_d2 = c["y_hat"] + 1, c["y_hat"] + 2
    c_d1 = [c_d2 + 1 + j for j in range(n_ocu)]

    _grupo(ws, 7, c_e, c_d1[-1], "Retropropagación")
    _encabezados(ws, 8, ["e = ŷ − y", "δ_salida"] + [f"δ_h{j + 1}" for j in range(n_ocu)], col=c_e)
    for k in range(len(X)):
        r = filas[0] + k
        yh = f"{_col(c['y_hat'])}{r}"
        _celda(ws, r, c_e, f"={yh}-{_col(c['y'])}{r}", formato=FMT_PESO)
        # δ_salida = ∂E/∂z_salida = 2·e/n · ŷ(1 − ŷ)
        _celda(ws, r, c_d2, f"=2*{_col(c_e)}{r}/{n}*{yh}*(1-{yh})", formato=FMT_PESO)
        for j in range(n_ocu):
            a = f"{_col(c['a'][j])}{r}"
            # δ_hj = δ_salida · w2_j · a_hj(1 − a_hj)
            _celda(ws, r, c_d1[j], f"={_col(c_d2)}{r}*{refs_ini['W2'][j]}*{a}*(1-{a})", formato=FMT_PESO)

    r = filas[1] + 2
    _celda(ws, r, 1, "ECM inicial (Excel)", negrita=True)
    _celda(ws, r, 3, f"=SUMXMY2({_rango(c['y_hat'], filas)},{_rango(c['y'], filas)})/{n}", formato=FMT_DIF)

    # Tabla de actualización: w_nuevo = w_inicial − η · ∂E/∂w
    r += 2
    _celda(ws, r, 1, "Actualización de parámetros: w ← w − η · ∂E/∂w", negrita=True)
    _encabezados(ws, r + 1, ["Parámetro", "", "Inicial", "∂E/∂w", "Nuevo (Excel)",
                             "Python (1 época)", "|Diferencia|"])
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
    parametros = []
    for i in range(n_ent):
        for j in range(n_ocu):
            parametros.append((f"W1[x{i + 1}, h{j + 1}]", refs_ini["W1"][i][j],
                               f"=SUMPRODUCT({_rango(c['x'][i], filas)},{_rango(c_d1[j], filas)})",
                               pesos_1["W1"][i, j]))
    for j in range(n_ocu):
        parametros.append((f"b1[h{j + 1}]", refs_ini["b1"][j],
                           f"=SUM({_rango(c_d1[j], filas)})", pesos_1["b1"][0, j]))
    for j in range(n_ocu):
        parametros.append((f"W2[h{j + 1}, ŷ]", refs_ini["W2"][j],
                           f"=SUMPRODUCT({_rango(c['a'][j], filas)},{_rango(c_d2, filas)})",
                           pesos_1["W2"][j, 0]))
    parametros.append(("b2", refs_ini["b2"], f"=SUM({_rango(c_d2, filas)})", pesos_1["b2"][0, 0]))

    inicio = r + 2
    for k, (nombre, ref_ini, gradiente, valor_py) in enumerate(parametros):
        f = inicio + k
        _celda(ws, f, 1, nombre, negrita=True)
        _celda(ws, f, 3, f"={ref_ini}", VERDE, formato=FMT_PESO)
        _celda(ws, f, 4, gradiente, formato=FMT_PESO)
        _celda(ws, f, 5, f"=C{f}-{eta}*D{f}", formato=FMT_PESO)
        _celda(ws, f, 6, float(valor_py), AZUL, formato=FMT_PESO)
        _celda(ws, f, 7, f"=ABS(E{f}-F{f})", formato=FMT_DIF)
    fin = inicio + len(parametros) - 1
    _celda(ws, fin + 2, 1, "Diferencia máxima Excel vs Python", negrita=True)
    _celda(ws, fin + 2, 7, f"=MAX(G{inicio}:G{fin})", formato=FMT_DIF)
    _leyenda(ws, fin + 4)
    _anchos(ws, [16, 10, 15, 15, 15, 18, 15] + [12] * (c_d1[-1] - 7))


def _hoja_iniciales(ws, pesos_iniciales, semilla):
    _titulo(ws, "Pesos iniciales",
            f"Valores aleatorios uniformes en [−1, 1] generados con la semilla {semilla}. "
            "La hoja Retropropagación parte de aquí.")
    refs = _escribir_pesos(ws, 4, pesos_iniciales)
    _anchos(ws, [26, 12] + [12] * (pesos_iniciales["W1"].shape[1] - 1))
    return refs


def _hoja_entrenamiento(ws, historial, cada):
    _titulo(ws, "Curva de entrenamiento",
            f"Error cuadrático medio cada {cada} épocas (época 0 = pesos iniciales).")
    _encabezados(ws, 4, ["Época", "ECM"])
    epocas = list(range(0, len(historial), cada))
    if epocas[-1] != len(historial) - 1:
        epocas.append(len(historial) - 1)
    for k, e in enumerate(epocas):
        _celda(ws, 5 + k, 1, e, AZUL)
        _celda(ws, 5 + k, 2, historial[e], AZUL, formato="0.000000")

    grafica = ScatterChart()
    grafica.title = "ECM vs época"
    grafica.x_axis.title = "Época"
    grafica.y_axis.title = "ECM"
    grafica.legend = None
    grafica.height, grafica.width = 11, 22
    serie = Series(Reference(ws, min_col=2, min_row=4, max_row=4 + len(epocas)),
                   Reference(ws, min_col=1, min_row=5, max_row=4 + len(epocas)), title_from_data=True)
    serie.marker.symbol = "none"
    serie.smooth = False
    serie.graphicalProperties.line.solidFill = "1F4E78"
    serie.graphicalProperties.line.width = 25000
    grafica.series.append(serie)
    grafica.x_axis.scaling.min, grafica.x_axis.scaling.max = 0, epocas[-1]
    grafica.y_axis.scaling.min = 0
    grafica.y_axis.numFmt = "0.00"
    grafica.x_axis.majorGridlines = None
    grafica.x_axis.delete = False
    grafica.y_axis.delete = False
    for titulo in (grafica.title, grafica.x_axis.title, grafica.y_axis.title):
        titulo.overlay = False  # que los títulos no se encimen sobre el área de la gráfica
    ws.add_chart(grafica, "D4")
    _anchos(ws, [10, 12])
    ws.freeze_panes = "A5"


# --------------------------------------------------------------------------- API

def exportar(archivo, config, X, y, pesos_iniciales, pesos_1, pesos_finales, historial, y_pred,
             cada=100):
    wb = Workbook()
    ws_pesos = wb.active
    ws_pesos.title = "Pesos"
    ws_comp = wb.create_sheet("Comprobación")
    ws_retro = wb.create_sheet("Retropropagación")
    ws_ini = wb.create_sheet("Pesos iniciales")
    ws_ent = wb.create_sheet("Entrenamiento")

    refs, celdas = _hoja_pesos(ws_pesos, config, pesos_finales, historial[-1])
    _hoja_comprobacion(ws_comp, X, y, refs, y_pred, celdas)
    refs_ini = _hoja_iniciales(ws_ini, pesos_iniciales, config["semilla"])
    _hoja_retropropagacion(ws_retro, X, y, refs_ini, pesos_1, celdas)
    _hoja_entrenamiento(ws_ent, historial, cada)

    wb.calculation.fullCalcOnLoad = True  # Excel evalúa todas las fórmulas al abrir
    wb.save(archivo)


def leer_pesos(archivo=ARCHIVO_PESOS):
    """Lee la configuración y los pesos de la hoja Pesos. Regresa (config, pesos)."""
    ws = load_workbook(archivo, data_only=True)["Pesos"]
    config = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
              for r in range(FILA_CONFIG, FILA_PESOS)}
    n_ent, n_ocu, _ = (int(v) for v in config["Arquitectura"].split("-"))

    def valor(fila, col):
        return float(ws.cell(row=fila, column=col).value)

    f = FILA_PESOS
    W1 = np.array([[valor(f + 2 + i, 2 + j) for j in range(n_ocu)] for i in range(n_ent)])
    b1 = np.array([[valor(f + 2 + n_ent, 2 + j) for j in range(n_ocu)]])
    f += 4 + n_ent
    W2 = np.array([[valor(f + 2 + j, 2)] for j in range(n_ocu)])
    b2 = np.array([[valor(f + 2 + n_ocu, 2)]])
    return config, {"W1": W1, "b1": b1, "W2": W2, "b2": b2}
